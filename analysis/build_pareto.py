"""Pareto (top 5 models by units and value) and 5-tm class grid (1 best-selling model per class).
Usage: python3 analysis/build_pareto.py <Crane_Market_GVW_Pairing_v11_Noir.xlsx>"""
import sys
import pandas as pd, numpy as np, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter
d=pd.read_excel(sys.argv[1],sheet_name='8. Pivot Source')
d=d[d.Brand!='ZOOMLION']
K='Kategori Max Lifting Moment';T='Max Lifting Moment asli (tm)';FX=17803
fam={'SPK32080B':'SPK32080','SPK32080C':'SPK32080','SPK36080C':'SPK36080','SPK42502C':'SPK42502','SPK42502E':'SPK42502',
     'PK 41002 EH C':'PK 41002','PK 41002 MH C':'PK 41002','PK 53002 SH B':'PK 53002','PK 53002 SH C':'PK 53002'}
d['m']=d.apply(lambda r: r.Model.replace(r.Brand+' ',''),axis=1); d['m']=d.m.map(lambda x: fam.get(x,x))
BR={'SANY PALFINGER':'Sany Palfinger','PALFINGER':'Palfinger','XCMG':'XCMG','AMCO VEBA':'Amco Veba','HIAB':'Hiab','HYVA':'Hyva','FASSI':'Fassi'}
rp=lambda v: round(v*FX/1e7)*1e7
band=lambda t: int(np.ceil(t/5)*5)
cat=json.load(open('data/ferrari_catalogue.json'))['models']; pl=json.load(open('data/ferrari_price_list.json'))['models']
price={}
for r in pl:
    k=r.get('catalogue_model'); 
    if k: price.setdefault(k,[]).append((r['ferrari_model'],r['price_crane_only_idr']))
wb=Workbook(); ws=wb.active; ws.title='Pareto'
H=Font(bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1F2A36'); RED=PatternFill('solid',fgColor='FBE3E5')
thin=Side(style='thin',color='C9D0D6'); B=Border(top=thin,bottom=thin,left=thin,right=thin)
row=1; popular={}
summary={}
for c,col in [('Medium','F2A33A'),('Heavy','4F9CF9')]:
    s=d[d[K]==c]; g=s.groupby(['Brand','m']).agg(tm=(T,'max'),q=('Quantity','sum'),v=('Total Value (USD)','sum')).reset_index()
    popular[c]=set()
    for key,lab in [('q','Units'),('v','Import value (IDR, before tax)')]:
        h=g.sort_values([key,'v' if key=='q' else 'q'],ascending=False).reset_index(drop=True)
        tot=h[key].sum(); h['share']=h[key]/tot; h['cum']=h.share.cumsum(); n80=int((h.cum<0.8).sum())+1
        top=h.head(5); popular[c]|=set(zip(top.Brand,top.m))
        summary[(c,key)]=(round(top.share.sum()*100),n80,len(h))
        ws.cell(row,1,f'{c} class: top 5 models by {lab.lower()}, Jan 2023 – 14 Aug 2026').font=Font(bold=True,size=12,color=col)
        ws.cell(row+1,1,f'Top 5 = {top.share.sum():.0%} of class {lab.split(" (")[0].lower()}; {n80} of {len(h)} models reach 80%. Variants merged (e.g. SPK32080/B/C, PK 41002 EH C/MH C).').font=Font(italic=True,color='5A6570')
        hdr=['#','Brand','Model','Max Lifting Moment (tm)','tm class','Units','Import value (IDR, before tax)','Share','Cumulative']
        for j,t in enumerate(hdr,1):
            x=ws.cell(row+2,j,t); x.font=H; x.fill=HF; x.alignment=Alignment(wrap_text=True,vertical='center'); x.border=B
        r0=row+3
        for i,rr in top.iterrows():
            vals=[i+1,BR[rr.Brand],rr.m,rr.tm,f'>{band(rr.tm)-5}–{band(rr.tm)}',int(rr.q),rp(rr.v),rr.share,rr.cum]
            for j,v in enumerate(vals,1):
                x=ws.cell(r0+i,j,v); x.border=B
                if j==7: x.number_format='"Rp "#,##0'
                if j in (8,9): x.number_format='0%'
        r5=r0+5
        oth=h.iloc[5:]; vals=['','Others',f'{len(oth)} models','','',int(oth.q.sum()),rp(oth.v.sum()),oth.share.sum(),1.0]
        for j,v in enumerate(vals,1):
            x=ws.cell(r5,j,v); x.border=B; x.font=Font(color='5A6570')
            if j==7: x.number_format='"Rp "#,##0'
            if j in (8,9): x.number_format='0%'
        ch=BarChart(); ch.title=f'{c}: top 5 models by {lab.split(" (")[0].lower()}'; ch.y_axis.title=lab.split(' (')[0]
        dc=6 if key=='q' else 7
        ch.add_data(Reference(ws,min_col=dc,min_row=row+2,max_row=r0+4),titles_from_data=True); ch.set_categories(Reference(ws,min_col=3,min_row=r0,max_row=r0+4))
        ch.series[0].graphicalProperties.solidFill=col
        ln=LineChart(); ln.add_data(Reference(ws,min_col=9,min_row=row+2,max_row=r0+4),titles_from_data=True); ln.y_axis.axId=200; ln.y_axis.scaling.max=1; ln.y_axis.number_format='0%'; ln.y_axis.crosses='max'; ln.y_axis.title='Cumulative'
        ch+=ln; ch.height=6.5; ch.width=14; ws.add_chart(ch,f'K{row}')
        row=r5+6
widths=[4,16,16,12,10,8,22,8,11]
for j,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(j)].width=w
# classification grid: one best-selling model (by units, then value) per 5-tm class
ORIGIN={'SANY PALFINGER':'china','XCMG':'china','PALFINGER':'eu','AMCO VEBA':'eu','HIAB':'eu','HYVA':'eu','FASSI':'eu'}
DUTY={'china':0,'eu':0.05}
taxed=lambda usd,b: usd*FX*(1+DUTY[ORIGIN[b]])*(1+0.11+0.025)
fmt=lambda v: 'Rp '+f'{round(v/1e7)*1e7:,.0f}'.replace(',','.')
def pick_variant(lst):
    names=[a for a,_ in lst]
    for r in pl:
        if r.get('use_in_deck') and r['ferrari_model'] in names: return next(x for x in lst if x[0]==r['ferrari_model'])
    for key in ('R A4',' A4'):
        for x in lst:
            if key in x[0]: return x
    return lst[0]
def ferrari_pick(lo,hi,ref_tm):
    fer=[m for m in cat if lo<m['max_lifting_moment_tm']<=hi]
    if not fer: return 'No Ferrari model', None
    priced=[m for m in fer if price.get(m['model'].replace(' ',''))]
    pool=priced or fer
    m=min(pool,key=lambda m: abs(m['max_lifting_moment_tm']-(ref_tm if ref_tm is not None else (lo+hi)/2)))
    if m in priced:
        a,p=pick_variant(price[m['model'].replace(' ','')])
        return f"{a} ({m['max_lifting_moment_tm']:g} tm) · {fmt(p)}", p
    return f"{m['model']} ({m['max_lifting_moment_tm']:g} tm) · price not yet known", None
gs=wb.create_sheet('Klasifikasi tm')
r=1; grid=[]
for c,col,bands in [('Medium','F2A33A',[30,35,40,45]),('Heavy','4F9CF9',[50,55,60,65,70,75,80,85,90])]:
    s=d[d[K]==c]; g=s.groupby(['Brand','m']).agg(tm=(T,'max'),q=('Quantity','sum'),v=('Total Value (USD)','sum')).reset_index()
    rows=[]
    for hi in bands:
        lo=hi-5; b=g[(g.tm>lo)&(g.tm<=hi)].sort_values(['q','v'],ascending=False)
        tot=int(b.q.sum()); top=b.iloc[0] if len(b) else None
        ftxt,fp=ferrari_pick(lo,hi,top.tm if top is not None else None)
        if top is None and ftxt=='No Ferrari model': continue
        rows.append((lo,hi,tot,top,ftxt,fp))
    brands=[b for b in ['SANY PALFINGER','PALFINGER','XCMG','AMCO VEBA','HIAB','HYVA','FASSI'] if any(x[3] is not None and x[3].Brand==b for x in rows)]
    gs.cell(r,1,f'{c} class: best-selling model in each 5-tm class vs F.lli Ferrari, Jan 2023 – 14 Aug 2026').font=Font(bold=True,size=12,color=col)
    gs.cell(r+1,1,'Cell = model (tm) · units · share of class units · price per unit, crane only (IDR, incl. tax). Best seller = most units in the class (variants merged).').font=Font(italic=True,color='5A6570')
    hdr=['tm class','Units in class']+[BR[b] for b in brands]+['F.lli Ferrari']
    for j,t in enumerate(hdr,1):
        x=gs.cell(r+2,j,t); x.font=H; x.fill=PatternFill('solid',fgColor='E30613' if t=='F.lli Ferrari' else '1F2A36'); x.border=B; x.alignment=Alignment(horizontal='center',wrap_text=True)
    rr=r+3; nc=len(hdr)
    for lo,hi,tot,top,ftxt,fp in rows:
        gs.cell(rr,1,f'>{lo}–{hi}').font=Font(bold=True); gs.cell(rr,2,tot if tot else 'No imports')
        if top is not None:
            up=taxed(top.v/top.q,top.Brand)
            x=gs.cell(rr,3+brands.index(top.Brand),f'{top.m} ({top.tm:g} tm)\n{int(top.q)} units · {top.q/tot:.0%} of class\n{fmt(up)} per unit'); x.font=Font(bold=True)
        fm=ftxt.split(' · ')[0] if ftxt!='No Ferrari model' else None
        grid.append(dict(cls=c,band=f'>{lo}–{hi}',units=tot,brand=top.Brand if top is not None else None,model=top.m if top is not None else None,
            tm=float(top.tm) if top is not None else None,top_units=int(top.q) if top is not None else 0,share=float(top.q/tot) if tot else 0,
            price_idr=float(taxed(top.v/top.q,top.Brand)) if top is not None else None,
            ferrari=fm.split(' (')[0] if fm else None,ferrari_tm=float(fm.split('(')[1].split(' ')[0]) if fm else None,ferrari_price_idr=fp))
        x=gs.cell(rr,nc,ftxt); x.fill=RED
        for j in range(1,nc+1):
            x=gs.cell(rr,j); x.border=B; x.alignment=Alignment(wrap_text=True,vertical='top',horizontal='center' if j<nc else 'left')
        gs.row_dimensions[rr].height=48
        rr+=1
    r=rr+2
gs.cell(r,1,'Competitor price = average import unit price 2023–2026 × Rp 17.803/USD (JISDOR 23 Sep 2026) × (1 + import duty: 0% China – ACFTA, 5% Europe – assumed) × (1 + PPN 11% + PPh 22 2.5%). Import cost before distributor margin.').font=Font(italic=True,color='5A6570',size=9)
gs.cell(r+1,1,'Ferrari price = TSP selling price to customers, crane only, incl. GP, warranty and tax (price list). Ferrari model = the priced catalogue model closest in tm to the best seller. Prices rounded to Rp 10.000.000.').font=Font(italic=True,color='5A6570',size=9)
gs.column_dimensions['A'].width=10; gs.column_dimensions['B'].width=9
for j in range(3,12): gs.column_dimensions[get_column_letter(j)].width=24
wb.save('analysis/Pareto_tm_Class.xlsx')
json.dump(grid,open('analysis/tm_grid.json','w'),indent=1)
for x in grid: print(x)
