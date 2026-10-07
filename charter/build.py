# Project charter (one slide) on the Triatra NEDP template. Run: python3 build.py NEDP_Background_Template.pptx Project_Charter_Flli_Ferrari.pptx
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
OR, TXT, MUT, CARD, LINE = 'FF5A1F', 'F4F1EE', 'A79F98', '171310', '3A322D'
FONT = 'Bahnschrift'
p = Presentation(sys.argv[1]); s = p.slides[0]
rgb = lambda h: RGBColor.from_string(h)
for sh in s.shapes:
    if not sh.has_text_frame: continue
    t = sh.text_frame.text
    new = {'SECTION 00 · SUBTITLE':'PROJECT CHARTER', 'Slide Title':'F.lli Ferrari Market Strategy',
           'NEDP MID YEAR REVIEW 2026':'F.LLI FERRARI MARKET STRATEGY · PROJECT CHARTER', '00 / 00':'01 / 01'}.get(t)
    if new: sh.text_frame.paragraphs[0].runs[0].text = new
    if t == 'Slide Title': sh.text_frame.paragraphs[0].runs[0].font.name = FONT; sh.text_frame.paragraphs[0].runs[0].font.bold = True

def box(x, y, w, h, fill=None, line=None):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill: r.fill.solid(); r.fill.fore_color.rgb = rgb(fill)
    else: r.fill.background()
    if line: r.line.color.rgb = rgb(line); r.line.width = Pt(0.75)
    else: r.line.fill.background()
    r.shadow.inherit = False
    return r
def text(x, y, w, h, paras, size=9, color=TXT, bold=False, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, ml=0.1):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(ml); tf.margin_top = tf.margin_bottom = Inches(0.03); tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        pg = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); pg.alignment = align; pg.space_after = Pt(1)
        for seg in (para if isinstance(para, list) else [para]):
            t, b = (seg, bold) if isinstance(seg, str) else seg
            r = pg.add_run(); r.text = t; f = r.font; f.name = FONT; f.size = Pt(size); f.bold = b; f.color.rgb = rgb(color)
    return tb
def head(x, y, w, t, sub=None, h=0.25):
    box(x, y, w, h + (0.18 if sub else 0), OR)
    text(x, y, w, h, [t], 10.5, 'FFFFFF', True, MSO_ANCHOR.MIDDLE)
    if sub: text(x, y + 0.2, w, 0.22, [sub], 9, 'FFFFFF', False, MSO_ANCHOR.MIDDLE)
    return y + h + (0.18 if sub else 0)
def body(x, y, w, h, paras, **k):
    box(x, y, w, h, CARD); text(x, y + 0.03, w, h - 0.03, paras, **k); return y + h + 0.06
num = lambda items: [f'{i}.  {t}' for i, t in enumerate(items, 1)]

L, LW, R, RW, Y0 = 0.46, 6.25, 6.86, 6.01, 1.5
# ---------- left column ----------
y = head(L, Y0, LW, 'PROJECT NAME')
y = body(L, y, LW, 0.27, [('F.lli Ferrari Market Strategy', True)], size=10)
y = head(L, y, LW, 'BACKGROUND')
y = body(L, y, LW, 1.02, [[ 'Triatra distributes F.lli Ferrari truck-mounted knuckle boom cranes in Indonesia. From Jan 2023 to Aug 2026, Indonesia imported 638 knuckle boom cranes; ',
    ('Medium and Heavy cranes (above 25 tm) are 41% of units but 64% of import value', True), ', led by Sany Palfinger and Palfinger. ',
    'F.lli Ferrari is cheaper than European rivals in five tm classes, but costs more than Chinese brands. ',
    ('This project is initiated to give a structured assessment of where F.lli Ferrari can compete, which cranes to stock and how to position them', True),
    ', so the F.lli Ferrari line grows on a sound, data-based footing.']], size=9)
y = head(L, y, LW, 'OBJECTIVES')
y = body(L, y, LW, 0.58, num(['Map Indonesia\'s knuckle boom crane market by Max Lifting Moment (tm) class, brand and model',
    'Compare F.lli Ferrari head to head with every competitor model on price, size and features',
    'Define which F.lli Ferrari cranes Triatra should stock and how to position them']), size=9)
y = head(L, y, LW, 'PROJECT SCOPE', 'F.LLI FERRARI CRANES, MEDIUM AND HEAVY CLASS (> 25 tm)')
box(L, y, LW, 0.55, CARD)
text(L, y + 0.03, LW / 2, 0.6, num(['Market size & brand leaders', 'Price comparison by 5-tm class', 'F.lli Ferrari crane advantage']), size=9)
text(L + LW / 2, y + 0.03, LW / 2, 0.6, ['4.  Stock & positioning strategy', '5.  Implementation timeline & sales plan'], size=9)
y += 0.61
kw = 4.6
head(L, y, kw, 'KPI / METRIC'); y = head(L + kw + 0.04, y, LW - kw - 0.04, 'TARGET')
box(L, y, kw, 0.88, CARD); box(L + kw + 0.04, y, LW - kw - 0.04, 0.88, CARD)
text(L, y + 0.03, kw, 0.92, num(['Market map by tm class, brand and model (2023 – 2026)', 'Competitor price and specification benchmark',
    'List of F.lli Ferrari cranes that win head to head', 'Stock recommendation per tm class', 'Implementation timeline and sales plan']), size=9)
text(L + kw + 0.04, y + 0.03, LW - kw - 0.04, 0.92, ['Delivered the F.lli Ferrari market strategy in March 2027'], size=9)
y += 0.94
y = head(L, y, LW, 'SUPPORT REQUIRED')
body(L, y, LW, 0.3, ['Marketing and Sales; United Tractors Pandu Engineering (UTPE); F.lli Ferrari Italia'], size=9)

# ---------- right column ----------
y = head(R, Y0, RW, 'PROJECT TIMEFRAME')
months = ['Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec', 'Jan', 'Feb', 'Mar']
acts = [('Planning – Framework', range(0, 6)), ('Observation, Study, and Review', range(4, 6)),
        ('Analysis and Formulation', range(6, 9)), ('Reporting and Evaluation', range(9, 10))]
aw = 2.05; mw = (RW - aw) / len(months); rh = 0.27; y += 0.04
box(R, y, aw, rh * 2, CARD, LINE); text(R, y, aw, rh * 2, [('Activity', True)], 9.5, anchor=MSO_ANCHOR.MIDDLE)
box(R + aw, y, mw * 7, rh, CARD, LINE); text(R + aw, y, mw * 7, rh, [('2026', True)], 9, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
box(R + aw + mw * 7, y, mw * 3, rh, CARD, LINE); text(R + aw + mw * 7, y, mw * 3, rh, [('2027', True)], 9, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
for i, m in enumerate(months):
    box(R + aw + i * mw, y + rh, mw, rh, CARD, LINE); text(R + aw + i * mw, y + rh, mw, rh, [m], 8.5, MUT, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, ml=0)
y += rh * 2
for a, rng in acts:
    box(R, y, aw, 0.3, None, LINE); text(R, y, aw, 0.3, [a], 9, anchor=MSO_ANCHOR.MIDDLE)
    for i in range(len(months)): box(R + aw + i * mw, y, mw, 0.3, OR if i in rng else None, LINE)
    y += 0.3
y += 0.08
y = head(R, y, RW, 'BENEFIT OF PROJECT')
y = body(R, y, RW, 0.45, ['Give Triatra a clear, data-based basis to choose which F.lli Ferrari cranes to stock, and to target the tm classes and customers where F.lli Ferrari wins on price and quality.'], size=9)
y = head(R, y, RW, 'PROJECT TEAM')
team = [('Executive Sponsor', 'Ricardo Tanada'), ('Advisor', 'Ricardo Tanada'), ('Coordinator', 'Ignatius Indrawan'),
        ('Project Lead', 'Tiberias Krisgaharu Simu'), ('Team Member', 'Yohannes Eurico Christanto, Muhammad Fawwaz Jauharin')]
box(R, y, RW, 0.92, CARD)
text(R, y + 0.03, 1.5, 0.9, [t for t, _ in team], size=9)
text(R + 1.45, y + 0.03, RW - 1.45, 0.9, [': ' + n for _, n in team], size=9)
y += 0.98
y = head(R, y, RW, 'PROJECT APPROVAL')
sig = [('Lead', 'Tiberias Krisgaharu Simu'), ('Coordinator', 'Ignatius Indrawan'), ('Advisor', 'Ricardo Tanada'), ('Executive Sponsor', 'Ricardo Tanada')]
g = 0.08; sw = (RW - g * 3) / 4; bottom = 7.02
for i, (r, n) in enumerate(sig):
    x = R + i * (sw + g)
    text(x, y + 0.02, sw, 0.36, [(r, True), 'Date:'], 8.5, ml=0.05)
    box(x, y + 0.4, sw, bottom - y - 0.4, None, LINE)
    box(x, bottom - 0.3, sw, 0.3, CARD, LINE); text(x, bottom - 0.3, sw, 0.3, [n], 8, anchor=MSO_ANCHOR.MIDDLE, ml=0.05)
p.save(sys.argv[2])
