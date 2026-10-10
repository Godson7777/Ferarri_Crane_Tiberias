"""Convert text baked into slide background images to native PowerPoint text; replace PCR with CRM.
usage: python3 fix.py orig.pptx out.pptx"""
import sys, io, copy
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
PX = 192.0                                   # background images are 2560 px wide = 13.333 in
FONT = 'Arial'
src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(src)

def img_of(slide):
    for sh in slide.shapes:
        if sh.shape_type == 13 and sh.width > Inches(13): return sh
def load(sh): return np.asarray(Image.open(io.BytesIO(sh.image.blob)).convert('RGB')).astype(float)
def save(sh, a):
    b = io.BytesIO(); Image.fromarray(np.clip(a, 0, 255).astype('uint8')).save(b, 'JPEG', quality=92)
    sh._element  # replace the image part blob in place (shared parts update every slide that uses it)
    part = sh.part.related_part(sh._element.blipFill.blip.rEmbed)
    part._blob = b.getvalue()

def text_bbox(a, x0, y0, x1, y1, thr=70):
    reg = a[y0:y1, x0:x1]; bg = np.median(reg.reshape(-1, 3), 0)
    m = np.abs(reg - bg).sum(2) > thr; ys, xs = np.where(m)
    return (x0 + xs.min(), y0 + ys.min(), x0 + xs.max() + 1, y0 + ys.max() + 1), bg
def fill_rect(a, x0, y0, x1, y1, col): a[y0:y1, x0:x1] = col
def fill_disc(a, cx, cy, r, col):
    yy, xx = np.ogrid[:a.shape[0], :a.shape[1]]; a[(yy - cy) ** 2 + (xx - cx) ** 2 <= r * r] = col

def add_text(slide, x0, y0, x1, y1, text, size, color, bold=True, align=PP_ALIGN.CENTER, spacing=None, alpha=None, wrap=False):
    tb = slide.shapes.add_textbox(Emu(int(x0 / PX * 914400)), Emu(int(y0 / PX * 914400)), Emu(int((x1 - x0) / PX * 914400)), Emu(int((y1 - y0) / PX * 914400)))
    tf = tb.text_frame; tf.word_wrap = wrap; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    p = tf.paragraphs[0]; p.alignment = align; r = p.add_run(); r.text = text
    f = r.font; f.name = FONT; f.size = Pt(size); f.bold = bold; f.color.rgb = RGBColor.from_string(color)
    rPr = r._r.get_or_add_rPr()
    if spacing is not None: rPr.set('spc', str(int(spacing * 100)))
    if alpha is not None:
        srgb = rPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
        a_el = srgb.makeelement(qn('a:alpha'), {'val': str(int(alpha * 100000))}); srgb.append(a_el)
    return tb
hexc = lambda c: '%02X%02X%02X' % tuple(int(v) for v in c)
pt_for = lambda cap_px: cap_px / PX * 72 / 0.716      # Arial cap/digit height = 0.716 em

S = lambda n: prs.slides[n - 1]
done = {}

# ---------- pills ("STEPS 1–5", "MAIN PROCESS · STEPS 1–3 / 4–5") and step badges ----------
pills = {13: ((2306, 214, 2456, 252), 'STEPS 1–5'), 14: ((2093, 214, 2456, 252), 'MAIN PROCESS · STEPS 1–3'), 15: ((2093, 214, 2456, 252), 'MAIN PROCESS · STEPS 4–5')}
badges = {13: [((118, 326, 173, 381), '1'), ((600, 330, 655, 385), '2'), ((1082, 330, 1137, 385), '3'), ((1562, 326, 1617, 381), '4'), ((2044, 330, 2099, 385), '5')],
          14: [((118, 326, 173, 381), '1'), ((920, 326, 975, 381), '2'), ((1724, 326, 1779, 381), '3')],
          15: [((118, 322, 173, 377), '4'), ((1324, 322, 1379, 377), '5')]}
for n in (13, 14, 15):
    sh = img_of(S(n)); a = load(sh)
    (x0, y0, x1, y1), t = pills[n]
    (bx0, by0, bx1, by1), bg = text_bbox(a, x0, y0, x1, y1)
    reg = a[by0:by1, bx0:bx1].reshape(-1, 3); col = reg[np.abs(reg - bg).sum(1).argmax()]
    fill_rect(a, x0, y0, x1, y1, bg)
    add_text(S(n), x0 - 10, y0, x1 + 10, y1, t, pt_for(by1 - by0) * 0.95, hexc(col), spacing=1.1)
    for (x0, y0, x1, y1), t in badges[n]:
        (tx0, ty0, tx1, ty1), bg = text_bbox(a, x0 + 8, y0 + 8, x1 - 8, y1 - 8)
        fill_rect(a, x0 + 8, y0 + 8, x1 - 7, y1 - 7, bg)
        add_text(S(n), x0, y0, x1, y1, t, pt_for(ty1 - ty0), 'FFFFFF')
    save(sh, a)

# ---------- slide 4 tags ----------
sh = img_of(S(4)); a = load(sh)
for (x0, y0, x1, y1), t in [((195, 979, 346, 1002), 'MY COMPANY'), ((2330, 1073, 2443, 1096), 'NEW · 2025')]:
    (tx0, ty0, tx1, ty1), bg = text_bbox(a, x0, y0, x1, y1)
    fill_rect(a, x0, y0, x1, y1, bg)
    add_text(S(4), x0 - 8, y0 - 3, x1 + 8, y1 + 3, t, pt_for(ty1 - ty0), 'FFFFFF', spacing=1.2)
save(sh, a)

# ---------- slide 6 legend, map markers, list markers ----------
sh = img_of(S(6)); a = load(sh)
from PIL import ImageFont
AR = ImageFont.truetype('/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf', 1000)
LEG = [(1913, 2022, 'Head Office'), (2089, 2222, 'Site Operation'), (2289, 2471, 'Site Representative')]   # ink x-extent, measured on the image
lsz = min((x1 - x0) / PX * 72 / (AR.getlength(t) / 1000) for x0, x1, t in LEG)
col = None
for x0, x1, t in LEG:
    (tx0, ty0, tx1, ty1), bg = text_bbox(a, x0 - 4, 296, x1 + 6, 330, 60)
    if col is None: reg = a[ty0:ty1, tx0:tx1].reshape(-1, 3); col = reg[np.abs(reg - bg).sum(1).argmax()]
    fill_rect(a, x0 - 4, 296, x1 + 6, 330, bg)
    add_text(S(6), x0 - 30, 292, x1 + 30, 334, t, lsz, hexc(col), bold=False, align=PP_ALIGN.CENTER)   # centred on the ink box (template lstStyle indents left-aligned text)
WHITE, ORANGE = {2, 7, 10, 14, 16}, {1}
mapm = {1: (817, 797), 2: (782, 773), 3: (667, 701), 4: (567, 551), 5: (1129, 687), 6: (1167, 701), 7: (1165, 645), 8: (1191, 633), 9: (1142, 603),
        10: (1225, 613), 11: (1182, 577), 12: (1191, 549), 13: (1253, 493), 14: (1256, 551), 15: (1370, 887), 16: (2081, 735)}
cols, rows = [104, 708, 1310, 1914], [999, 1095, 1193, 1289]
listm = {c * 4 + r + 1: (cols[c], rows[r]) for c in range(4) for r in range(4)}
for marks, R, size in [(mapm, 17, 7.5), (listm, 16, 6.5)]:
    for k, (cx, cy) in marks.items():
        # refine centre: centroid of non-background pixels in the window
        win = a[cy - R - 6:cy + R + 7, cx - R - 6:cx + R + 7]; bgw = np.median(win.reshape(-1, 3), 0)
        m = np.abs(win - bgw).sum(2) > 120; ys, xs = np.where(m)
        if len(xs): cx, cy = cx - R - 6 + (xs.min() + xs.max()) // 2, cy - R - 6 + (ys.min() + ys.max()) // 2
        yy, xx = np.ogrid[-R:R + 1, -R:R + 1]; disc = a[cy - R:cy + R + 1, cx - R:cx + R + 1][(yy ** 2 + xx ** 2) <= (0.62 * R) ** 2]
        col = np.median(disc, 0); fill_disc(a, cx, cy, 0.66 * R, col)
        tc = '111111' if k in WHITE else 'FFFFFF'
        add_text(S(6), cx - R, cy - R, cx + R + 1, cy + R + 1, str(k), size, tc)
save(sh, a)

# ---------- divider numbers (00, 02 … 06): rebuild the striped background, then a native translucent number ----------
from PIL import ImageFont, ImageDraw
AB = ImageFont.truetype('/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf', 1000)   # Arial Bold metrics
def gm(ch):                                        # ink metrics in em: left bearing, ink height, ink width, advance
    im = Image.new('L', (2000, 1600)); ImageDraw.Draw(im).text((300, 1200), ch, font=AB, fill=255, anchor='ls')
    x0, y0, x1, y1 = im.getbbox(); return (x0 - 300) / 1000, (y1 - y0) / 1000, (x1 - x0) / 1000, AB.getlength(ch) / 1000
divs = {2: '00', 5: '02', 7: '03', 9: '04', 12: '05', 16: '06'}
Y0, Y1, X0 = 192, 770, 1500
for n, num in divs.items():
    sh = img_of(S(n)); a = load(sh); orig = a.copy()
    top = a[Y0 - 10:Y0, X0:].mean(0); bot = a[Y1:Y1 + 10, X0:].mean(0)
    t = ((np.arange(Y0, Y1) - (Y0 - 5)) / ((Y1 + 5) - (Y0 - 5)))[:, None, None]
    est = top[None] * (1 - t) + bot[None] * t
    glyph = (orig[Y0:Y1, X0:] - est).sum(2) > 25; ys, xs = np.where(glyph)
    gx0, gx1, gy0, gy1 = X0 + xs.min(), X0 + xs.max() + 1, Y0 + ys.min(), Y0 + ys.max() + 1
    alpha = float(np.median(((orig[Y0:Y1, X0:] - est) / np.maximum(255 - est, 1))[glyph]))
    p0, p1 = max(gx0 - 14, X0) - X0, min(gx1 + 14, a.shape[1]) - X0
    a[Y0:Y1, X0 + p0:X0 + p1] = est[:, p0:p1]
    save(sh, a)
    l0, h0, w0, adv0 = gm(num[0]); l1, h1, w1, adv1 = gm(num[1])
    em = (gy1 - gy0) / max(h0, h1)                 # px per em
    spc = (gx1 - gx0) - (adv0 - l0 + l1 + w1) * em  # extra space between the two digits, px
    size = em / PX * 72
    bx = gx0 - l0 * em; cy = (gy0 + gy1) / 2 - 0.0605 * em        # offset measured against the original render
    for m in prs.slides:
        im = img_of(m)
        if im is not None and im.part.related_part(im._element.blipFill.blip.rEmbed) is sh.part.related_part(sh._element.blipFill.blip.rEmbed):
            add_text(m, bx, cy - 0.65 * em, bx + 1.6 * em, cy + 0.65 * em, num, size, 'FFFFFF', align=PP_ALIGN.LEFT, alpha=alpha, spacing=spc / PX * 72, wrap=True)
    done[n] = (int(gx0), int(gy0), int(gx1), int(gy1), round(alpha, 3), round(size), round(spc / PX * 72, 1))

# ---------- PCR -> CRM ----------
cnt = 0
for sl in prs.slides:
    for sh in sl.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if 'PCR' in r.text: cnt += r.text.count('PCR'); r.text = r.text.replace('PCR', 'CRM')
prs.save(dst)
print('dividers', done, 'PCR replaced', cnt)
